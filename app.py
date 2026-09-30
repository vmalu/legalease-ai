import io
import requests
import streamlit as st
from docx import Document
from fpdf import FPDF

st.set_page_config(page_title="LegalEase", layout="centered")

st.markdown("<h2 style='text-align: center;'>⚖️ LegalEase</h2>", unsafe_allow_html=True)
st.markdown("<h4 style='text-align: center;'>AI Legal Document Generator</h4>", unsafe_allow_html=True)

def sanitize_text(text: str) -> str:
    return (
        text.replace("’", "'")
        .replace("‘", "'")
        .replace("“", '"')
        .replace("”", '"')
        .replace("—", "-")
    )

class LegalDocPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 8, "LegalEase", align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, "LegalEase Inc | contact@legalease.com | All Rights Reserved", align="C")

def format_docx(text: str, doc_type: str) -> io.BytesIO:
    doc = Document()
    doc.add_heading(f"LegalEase - {doc_type}", level=0)
    for block in text.split("\n\n"):
        clean_block = block.strip()
        if clean_block:
            if clean_block.startswith("## "):
                doc.add_heading(clean_block.replace("## ", ""), level=1)
            elif clean_block in ["Between:", "And:", "WITNESSETH:"]:
                p = doc.add_paragraph()
                run = p.add_run(clean_block)
                run.bold = True
            else:
                doc.add_paragraph(clean_block)

    section = doc.sections[0]
    footer = section.footer
    footer.paragraphs[0].text = "LegalEase Inc | contact@legalease.com | All Rights Reserved"

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalDocPDF(orientation="P", unit="mm", format="A4")
    pdf.set_margins(left=20, top=20, right=20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    
    # Page-ன் நிகர அகலம் (A4 width 210mm - 40mm margins = 170mm)
    w = 170
    
    clean = sanitize_text(text).encode("latin-1", "replace").decode("latin-1")
    
    for paragraph in clean.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
            
        # எப்போதும் புதிய பத்தி தொடங்கும் போது cursor இடது மார்ஜினுக்கு வரவைக்க:
        pdf.set_x(20)
        
        if paragraph.startswith("## "):
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(w, 7, paragraph.replace("## ", ""), align="C")
            pdf.ln(8)
        elif paragraph in ["Between:", "And:", "WITNESSETH:"]:
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(w, 6, paragraph, align="L")
            pdf.ln(6)
        elif any(paragraph.startswith(f"{i}.") for i in range(1, 15)):
            parts = paragraph.split("\n", 1)
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_x(20)
            pdf.multi_cell(w, 5, parts[0].strip(), align="L")
            pdf.ln(1)
            if len(parts) > 1:
                pdf.set_font("Helvetica", "", 10)
                pdf.set_x(20)
                pdf.multi_cell(w, 5, parts[1].strip(), align="L")
            pdf.ln(4)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(w, 5, paragraph, align="L")
            pdf.ln(4)
            
    return bytes(pdf.output())

# Input Form
document_type = st.text_input("Document Type (Ex: Agreement, Contract, NDA)")
parties = st.text_area("Parties Involved")
terms = st.text_area("Terms & Conditions (Use semicolons for bullet points)")
dates = st.text_input("Effective Date")

if st.button("Generate Document"):
    if not (document_type and parties and terms and dates):
        st.warning("Please fill in all the fields.")
    else:
        with st.spinner("Generating document..."):
            try:
                response = requests.post(
                    "http://localhost:8000/generate",
                    json={
                        "document_type": document_type,
                        "parties": parties,
                        "terms": terms,
                        "dates": dates
                    }
                )
                if response.status_code == 200:
                    st.session_state["generated_text"] = sanitize_text(response.json()["document"])
                    st.session_state["doc_type"] = document_type
                    st.success("Document Generated Successfully!")
                else:
                    st.error("Failed to generate document. Ensure FastAPI server is running.")
            except Exception as e:
                st.error(f"Error connecting to backend: {e}")

# Output Section
if "generated_text" in st.session_state:
    st.markdown("---")
    edit_mode = st.checkbox("Click to Edit Document")
    
    if edit_mode:
        edited = st.text_area("Edit Document Below:", value=st.session_state["generated_text"], height=420)
        st.session_state["generated_text"] = edited
    else:
        st.markdown(st.session_state["generated_text"])

    st.markdown("---")
    curr_doc = st.session_state["generated_text"]
    curr_name = st.session_state.get("doc_type", "document").replace(" ", "_").lower()

    col1, col2, col3 = st.columns(3)
    with col1:
        st.download_button("Download as .TXT", data=curr_doc, file_name=f"{curr_name}.txt", mime="text/plain")
    with col2:
        st.download_button("Download as .DOCX", data=format_docx(curr_doc, st.session_state.get("doc_type", "Document")), file_name=f"{curr_name}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    with col3:
        st.download_button("Download as .PDF", data=format_pdf(curr_doc, st.session_state.get("doc_type", "Document")), file_name=f"{curr_name}.pdf", mime="application/pdf")